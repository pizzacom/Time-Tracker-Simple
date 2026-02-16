"""Export service - PDF, CSV, Excel generation."""
import io
import csv
import uuid
from datetime import date
from typing import List, Optional

from sqlalchemy import select, and_, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.time_entry import TimeEntry
from app.models.overtime import OvertimeLedger
from app.models.holiday import Holiday
from app.models.absence import Absence
from app.models.user import User
from app.services.work_schedule_service import get_soll_minutes_for_date
from app.utils.time_calc import (
    get_month_days, format_duration, get_weekday_name_de, time_to_minutes,
)


async def _get_month_data(
    db: AsyncSession, user_id: uuid.UUID, year: int, month: int
):
    """Gather all data needed for export."""
    days = get_month_days(year, month)

    # Get entries for the month
    result = await db.execute(
        select(TimeEntry)
        .where(
            and_(
                TimeEntry.user_id == user_id,
                TimeEntry.date >= days[0],
                TimeEntry.date <= days[-1],
            )
        )
        .order_by(TimeEntry.date, TimeEntry.start_time)
    )
    entries = result.scalars().all()

    # Get holidays
    holiday_result = await db.execute(
        select(Holiday).where(
            and_(Holiday.date >= days[0], Holiday.date <= days[-1])
        )
    )
    holidays = {h.date: h for h in holiday_result.scalars().all()}

    # Get absences
    absence_result = await db.execute(
        select(Absence).where(
            and_(
                Absence.user_id == user_id,
                Absence.date >= days[0],
                Absence.date <= days[-1],
            )
        )
    )
    absences = {a.date: a for a in absence_result.scalars().all()}

    # Build day-by-day data
    day_data = []
    total_soll = 0
    total_ist = 0
    for d in days:
        # Use raw soll (schedule-based, ignoring holidays/absences)
        raw_soll = await get_soll_minutes_for_date(db, user_id, d, raw=True)
        day_entries = [e for e in entries if e.date == d]
        ist = sum(e.work_minutes for e in day_entries)

        is_holiday = d in holidays
        absence = absences.get(d)

        # Credit holidays and absences as worked time equal to schedule soll
        credited = 0
        if is_holiday and raw_soll > 0:
            credited = raw_soll
        if absence and raw_soll > 0:
            credited = raw_soll

        ist += credited
        total_soll += raw_soll
        total_ist += ist

        day_data.append({
            "date": d,
            "weekday": get_weekday_name_de(d),
            "soll": raw_soll,
            "ist": ist,
            "delta": ist - raw_soll,
            "entries": day_entries,
            "holiday": holidays.get(d),
            "absence": absence,
            "credited": credited,
        })

    return day_data, total_soll, total_ist


async def generate_csv(
    db: AsyncSession, user_id: uuid.UUID, year: int, month: int
) -> bytes:
    """Generate CSV export for a month."""
    day_data, total_soll, total_ist = await _get_month_data(db, user_id, year, month)

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)

    # Header
    writer.writerow([
        "Datum", "Tag", "Von", "Bis", "Pause (min)",
        "Arbeit (min)", "Soll (min)", "Delta (min)", "Typ", "Überstunden", "Beschreibung"
    ])

    for day in day_data:
        if day["entries"]:
            for entry in day["entries"]:
                overtime_mark = ""
                if entry.entry_type == "overtime_used":
                    overtime_mark = "ÜS genommen"
                elif day["delta"] > 0 and len(day["entries"]) == 1:
                    overtime_mark = f"+{day['delta']} min"
                elif day["delta"] < 0 and len(day["entries"]) == 1:
                    overtime_mark = f"{day['delta']} min"
                writer.writerow([
                    day["date"].strftime("%d.%m.%Y"),
                    day["weekday"],
                    entry.start_time.strftime("%H:%M"),
                    entry.end_time.strftime("%H:%M"),
                    entry.break_minutes,
                    entry.work_minutes,
                    day["soll"],
                    entry.work_minutes - day["soll"] if len(day["entries"]) == 1 else "",
                    entry.entry_type,
                    overtime_mark,
                    entry.description or "",
                ])
        else:
            note = ""
            if day["holiday"]:
                note = f"Feiertag: {day['holiday'].name}"
            elif day["absence"]:
                note = f"Abwesend: {day['absence'].type}"
            writer.writerow([
                day["date"].strftime("%d.%m.%Y"),
                day["weekday"],
                "", "", "",
                day["credited"] if day["credited"] else "",
                day["soll"],
                day["delta"],
                "", "", note,
            ])

    # Summary row
    writer.writerow([])
    writer.writerow([
        "SUMME", "", "", "", "",
        total_ist, total_soll, total_ist - total_soll, "", ""
    ])

    csv_content = output.getvalue()
    # BOM for Excel UTF-8 compatibility
    return ("\ufeff" + csv_content).encode("utf-8")


async def generate_excel(
    db: AsyncSession, user_id: uuid.UUID, year: int, month: int
) -> bytes:
    """Generate Excel export for a month."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    # Get user info
    user_result = await db.execute(select(User).where(User.id == user_id))
    user = user_result.scalar_one()

    day_data, total_soll, total_ist = await _get_month_data(db, user_id, year, month)

    wb = Workbook()

    # Sheet 1: Zeiterfassung
    ws = wb.active
    ws.title = "Zeiterfassung"

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    alt_fill = PatternFill(start_color="FAFAFC", end_color="FAFAFC", fill_type="solid")
    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    bold_font = Font(bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    # Title
    ws.merge_cells("A1:J1")
    ws["A1"] = f"Zeiterfassung - {user.full_name}"
    ws["A1"].font = Font(bold=True, size=14)

    month_names = [
        "Januar", "Februar", "März", "April", "Mai", "Juni",
        "Juli", "August", "September", "Oktober", "November", "Dezember"
    ]
    ws.merge_cells("A2:J2")
    ws["A2"] = f"{month_names[month - 1]} {year}"
    ws["A2"].font = Font(size=11, color="666666")

    # Headers (row 4)
    headers = ["Datum", "Tag", "Von", "Bis", "Pause", "Arbeit", "Soll", "Delta", "Typ", "Überstunden", "Beschreibung"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Data rows
    row = 5
    for i, day in enumerate(day_data):
        if day["entries"]:
            for entry in day["entries"]:
                ws.cell(row=row, column=1, value=day["date"].strftime("%d.%m.%Y")).border = thin_border
                ws.cell(row=row, column=2, value=day["weekday"]).border = thin_border
                ws.cell(row=row, column=3, value=entry.start_time.strftime("%H:%M")).border = thin_border
                ws.cell(row=row, column=4, value=entry.end_time.strftime("%H:%M")).border = thin_border
                ws.cell(row=row, column=5, value=format_duration(entry.break_minutes)).border = thin_border
                ws.cell(row=row, column=6, value=format_duration(entry.work_minutes)).border = thin_border
                ws.cell(row=row, column=7, value=format_duration(day["soll"])).border = thin_border
                delta = day["delta"]
                delta_cell = ws.cell(row=row, column=8, value=format_duration(delta))
                delta_cell.border = thin_border
                if delta > 0:
                    delta_cell.fill = green_fill
                elif delta < 0:
                    delta_cell.fill = red_fill
                ws.cell(row=row, column=9, value=entry.entry_type).border = thin_border
                overtime_mark = ""
                if entry.entry_type == "overtime_used":
                    overtime_mark = "ÜS genommen"
                elif day["delta"] > 0:
                    overtime_mark = f"+{format_duration(day['delta'])}"
                elif day["delta"] < 0:
                    overtime_mark = format_duration(day["delta"])
                ot_cell = ws.cell(row=row, column=10, value=overtime_mark)
                ot_cell.border = thin_border
                if entry.entry_type == "overtime_used":
                    ot_cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                ws.cell(row=row, column=11, value=entry.description or "").border = thin_border

                if i % 2 == 0:
                    for c in range(1, 12):
                        if ws.cell(row=row, column=c).fill == PatternFill():
                            ws.cell(row=row, column=c).fill = alt_fill
                row += 1
        else:
            note = ""
            if day["holiday"]:
                note = f"Feiertag: {day['holiday'].name}"
            elif day["absence"]:
                note = f"{day['absence'].type}"

            ws.cell(row=row, column=1, value=day["date"].strftime("%d.%m.%Y")).border = thin_border
            ws.cell(row=row, column=2, value=day["weekday"]).border = thin_border
            ws.cell(row=row, column=3, value="–").border = thin_border
            ws.cell(row=row, column=4, value="–").border = thin_border
            ws.cell(row=row, column=5, value="–").border = thin_border
            credited_str = format_duration(day["credited"]) if day.get("credited") else "–"
            ws.cell(row=row, column=6, value=credited_str).border = thin_border
            ws.cell(row=row, column=7, value=format_duration(day["soll"])).border = thin_border
            delta_cell = ws.cell(row=row, column=8, value=format_duration(day["delta"]))
            delta_cell.border = thin_border
            if day["delta"] > 0:
                delta_cell.fill = green_fill
            elif day["delta"] < 0:
                delta_cell.fill = red_fill
            ws.cell(row=row, column=9, value="").border = thin_border
            ws.cell(row=row, column=10, value="").border = thin_border
            ws.cell(row=row, column=11, value=note).border = thin_border

            if i % 2 == 0:
                for c in range(1, 12):
                    ws.cell(row=row, column=c).fill = alt_fill
            row += 1

    # Summary row
    row += 1
    summary_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
    ws.cell(row=row, column=1, value="SUMME").font = bold_font
    ws.cell(row=row, column=6, value=format_duration(total_ist)).font = bold_font
    ws.cell(row=row, column=7, value=format_duration(total_soll)).font = bold_font
    ws.cell(row=row, column=8, value=format_duration(total_ist - total_soll)).font = bold_font
    for c in range(1, 12):
        ws.cell(row=row, column=c).fill = summary_fill
        ws.cell(row=row, column=c).border = thin_border

    # Column widths
    widths = [12, 6, 8, 8, 8, 8, 8, 8, 12, 14, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    # Freeze header
    ws.freeze_panes = "A5"

    # Sheet 2: Überstunden
    ws2 = wb.create_sheet("Überstunden")
    overtime_result = await db.execute(
        select(OvertimeLedger)
        .where(OvertimeLedger.user_id == user_id)
        .order_by(OvertimeLedger.date)
    )
    overtime_entries = overtime_result.scalars().all()

    ws2["A1"] = "Überstunden-Verlauf"
    ws2["A1"].font = Font(bold=True, size=14)

    ot_headers = ["Datum", "Grund", "Minuten", "Laufender Saldo", "Notiz"]
    for col, header in enumerate(ot_headers, 1):
        cell = ws2.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    running_balance = 0
    for i, ot in enumerate(overtime_entries, 4):
        running_balance += ot.minutes
        ws2.cell(row=i, column=1, value=ot.date.strftime("%d.%m.%Y")).border = thin_border
        ws2.cell(row=i, column=2, value=ot.reason).border = thin_border
        ws2.cell(row=i, column=3, value=ot.minutes).border = thin_border
        ws2.cell(row=i, column=4, value=running_balance).border = thin_border
        ws2.cell(row=i, column=5, value=ot.note or "").border = thin_border

    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 10
    ws2.column_dimensions["D"].width = 15
    ws2.column_dimensions["E"].width = 30

    # Save to bytes
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


async def generate_pdf(
    db: AsyncSession, user_id: uuid.UUID, year: int, month: int
) -> bytes:
    """Generate PDF export for a month using reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    user_result = await db.execute(select(User).where(User.id == user_id))
    user = user_result.scalar_one()

    day_data, total_soll, total_ist = await _get_month_data(db, user_id, year, month)

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=12*mm, rightMargin=12*mm,
                           topMargin=10*mm, bottomMargin=10*mm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Heading1"], fontSize=14,
        spaceAfter=2, textColor=colors.HexColor("#282828"),
    )
    subtitle_style = ParagraphStyle(
        "CustomSubtitle", parent=styles["Normal"], fontSize=9,
        spaceAfter=2, textColor=colors.HexColor("#555555"),
    )
    small_style = ParagraphStyle(
        "Small", parent=styles["Normal"], fontSize=6,
        textColor=colors.HexColor("#333333"), leading=7,
    )

    month_names = [
        "Januar", "Februar", "März", "April", "Mai", "Juni",
        "Juli", "August", "September", "Oktober", "November", "Dezember"
    ]

    elements = []
    elements.append(Paragraph("Zeiterfassung", title_style))
    elements.append(Paragraph(f"{month_names[month - 1]} {year} — {user.full_name}", subtitle_style))
    elements.append(Spacer(1, 3*mm))

    # Build table data — skip days with no entries and no soll (e.g. free weekends)
    table_data = [["Datum", "Tag", "Von", "Bis", "Pause", "Arbeit", "Soll", "Delta", "Beschr."]]

    for day in day_data:
        # Skip days that have no entries, no soll, no holiday, no absence
        if not day["entries"] and day["soll"] == 0 and not day["holiday"] and not day["absence"]:
            continue

        date_str = day["date"].strftime("%d.%m")
        weekday = day["weekday"]

        if day["entries"]:
            for entry in day["entries"]:
                table_data.append([
                    date_str, weekday,
                    entry.start_time.strftime("%H:%M"),
                    entry.end_time.strftime("%H:%M"),
                    format_duration(entry.break_minutes),
                    format_duration(entry.work_minutes),
                    format_duration(day["soll"]),
                    format_duration(day["delta"]),
                    Paragraph(entry.description or "–", small_style),
                ])
        else:
            note = ""
            if day["holiday"]:
                note = day["holiday"].name
            elif day["absence"]:
                note = day["absence"].type
            credited_str = format_duration(day["credited"]) if day.get("credited") else "–"
            table_data.append([
                date_str, weekday, "–", "–", "–", credited_str,
                format_duration(day["soll"]),
                format_duration(day["delta"]),
                Paragraph(note, small_style),
            ])

    # Summary
    table_data.append([
        "Summe", "", "", "", "",
        format_duration(total_ist),
        format_duration(total_soll),
        format_duration(total_ist - total_soll),
        "",
    ])

    col_widths = [16*mm, 10*mm, 13*mm, 13*mm, 13*mm, 15*mm, 13*mm, 13*mm, None]
    available_width = A4[0] - 24*mm
    used = sum(w for w in col_widths if w)
    col_widths[-1] = available_width - used

    table = Table(table_data, colWidths=col_widths, repeatRows=1)

    style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 6),
        ("FONTSIZE", (0, 1), (-1, -1), 6),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDDDDD")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#FAFAFC")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F5F5F5")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ])
    table.setStyle(style)

    elements.append(table)
    elements.append(Spacer(1, 3*mm))

    # Get overtime balance for summary
    from app.services.overtime_service import get_overtime_balance
    overtime_balance = await get_overtime_balance(db, user_id)

    # Summary line
    summary_text = (
        f"Arbeitszeit: {format_duration(total_ist)}  |  "
        f"Soll: {format_duration(total_soll)}  |  "
        f"Saldo Monat: {format_duration(total_ist - total_soll)}  |  "
        f"<b>Überstunden gesamt: {format_duration(overtime_balance)}</b>"
    )
    elements.append(Paragraph(summary_text, subtitle_style))

    # Generated date
    from datetime import datetime
    gen_text = f"Erstellt am {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    elements.append(Spacer(1, 2*mm))
    generated_style = ParagraphStyle(
        "Generated", parent=styles["Normal"], fontSize=7,
        textColor=colors.HexColor("#999999"),
    )
    elements.append(Paragraph(gen_text, generated_style))

    # Signature lines
    elements.append(Spacer(1, 10*mm))
    sig_data = [["____________________", "", "____________________"],
                ["Mitarbeiter", "", "Vorgesetzter"]]
    sig_table = Table(sig_data, colWidths=[60*mm, 40*mm, 60*mm])
    sig_table.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#999999")),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
    ]))
    elements.append(sig_table)

    doc.build(elements)
    output.seek(0)
    return output.read()


async def generate_department_excel(
    db: AsyncSession, dept_id: uuid.UUID, year: int, month: int
) -> bytes:
    """Generate department Excel with one sheet per worker + summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    # Get all users in department
    result = await db.execute(
        select(User).where(
            and_(User.department_id == dept_id, User.is_active == True)
        ).order_by(User.last_name)
    )
    users = result.scalars().all()

    wb = Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    bold_font = Font(bold=True, size=10)
    green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    month_names = [
        "Januar", "Februar", "März", "April", "Mai", "Juni",
        "Juli", "August", "September", "Oktober", "November", "Dezember"
    ]

    summary_data = []

    # Create a sheet per worker
    for idx, user in enumerate(users):
        day_data, total_soll, total_ist = await _get_month_data(db, user.id, year, month)

        sheet_name = f"{user.last_name} {user.first_name}"[:31]
        if idx == 0:
            ws = wb.active
            ws.title = sheet_name
        else:
            ws = wb.create_sheet(title=sheet_name)

        # Title
        ws.merge_cells("A1:H1")
        ws["A1"] = f"Zeiterfassung - {user.full_name}"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A2:H2")
        ws["A2"] = f"{month_names[month - 1]} {year}"
        ws["A2"].font = Font(size=11, color="666666")

        # Headers
        headers = ["Datum", "Tag", "Von", "Bis", "Pause", "Arbeit", "Soll", "Delta"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        row = 5
        for dd in day_data:
            d = dd["date"]
            entries = dd["entries"]
            if entries:
                for e in entries:
                    ws.cell(row=row, column=1, value=d.strftime("%d.%m.%Y")).border = thin_border
                    ws.cell(row=row, column=2, value=dd["weekday"]).border = thin_border
                    ws.cell(row=row, column=3, value=e.start_time.strftime("%H:%M") if e.start_time else "").border = thin_border
                    ws.cell(row=row, column=4, value=e.end_time.strftime("%H:%M") if e.end_time else "").border = thin_border
                    ws.cell(row=row, column=5, value=format_duration(e.break_minutes)).border = thin_border
                    ws.cell(row=row, column=6, value=format_duration(e.work_minutes)).border = thin_border
                    ws.cell(row=row, column=7, value=format_duration(dd["soll"])).border = thin_border
                    delta_cell = ws.cell(row=row, column=8, value=format_duration(dd["delta"]))
                    delta_cell.border = thin_border
                    if dd["delta"] > 0:
                        delta_cell.fill = green_fill
                    elif dd["delta"] < 0:
                        delta_cell.fill = red_fill
                    row += 1
            else:
                ws.cell(row=row, column=1, value=d.strftime("%d.%m.%Y")).border = thin_border
                ws.cell(row=row, column=2, value=dd["weekday"]).border = thin_border
                typ = ""
                if dd.get("holiday"):
                    typ = f"Feiertag: {dd['holiday'].name}"
                elif dd.get("absence"):
                    typ = dd["absence"].type
                credited_str = format_duration(dd["credited"]) if dd.get("credited") else ""
                ws.cell(row=row, column=6, value=credited_str or typ).border = thin_border
                ws.cell(row=row, column=7, value=format_duration(dd["soll"])).border = thin_border
                delta_cell = ws.cell(row=row, column=8, value=format_duration(dd["delta"]))
                delta_cell.border = thin_border
                row += 1

        # Totals
        row += 1
        ws.cell(row=row, column=6, value="Gesamt Ist:").font = bold_font
        ws.cell(row=row, column=7, value=format_duration(total_ist)).font = bold_font
        row += 1
        ws.cell(row=row, column=6, value="Gesamt Soll:").font = bold_font
        ws.cell(row=row, column=7, value=format_duration(total_soll)).font = bold_font
        row += 1
        ws.cell(row=row, column=6, value="Delta:").font = bold_font
        delta_cell = ws.cell(row=row, column=7, value=format_duration(total_ist - total_soll))
        delta_cell.font = bold_font

        # Set column widths
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 5
        for c in ["C", "D", "E", "F", "G", "H"]:
            ws.column_dimensions[c].width = 10

        # Get overtime balance
        balance_result = await db.execute(
            select(func.coalesce(func.sum(OvertimeLedger.minutes), 0)).where(
                OvertimeLedger.user_id == user.id
            )
        )
        balance = balance_result.scalar()

        summary_data.append({
            "name": user.full_name,
            "soll": total_soll,
            "ist": total_ist,
            "delta": total_ist - total_soll,
            "balance": balance,
        })

    # Summary sheet
    ws_sum = wb.create_sheet(title="Übersicht", index=0)
    ws_sum.merge_cells("A1:E1")
    ws_sum["A1"] = f"Abteilungsübersicht - {month_names[month - 1]} {year}"
    ws_sum["A1"].font = Font(bold=True, size=14)

    sum_headers = ["Mitarbeiter", "Soll", "Ist", "Delta", "Überstunden-Saldo"]
    for col, header in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=3, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    dept_soll = 0
    dept_ist = 0
    for row_idx, sd in enumerate(summary_data, 4):
        ws_sum.cell(row=row_idx, column=1, value=sd["name"]).border = thin_border
        ws_sum.cell(row=row_idx, column=2, value=format_duration(sd["soll"])).border = thin_border
        ws_sum.cell(row=row_idx, column=3, value=format_duration(sd["ist"])).border = thin_border
        delta_cell = ws_sum.cell(row=row_idx, column=4, value=format_duration(sd["delta"]))
        delta_cell.border = thin_border
        if sd["delta"] > 0:
            delta_cell.fill = green_fill
        elif sd["delta"] < 0:
            delta_cell.fill = red_fill
        bal_cell = ws_sum.cell(row=row_idx, column=5, value=format_duration(sd["balance"]))
        bal_cell.border = thin_border
        dept_soll += sd["soll"]
        dept_ist += sd["ist"]

    # Department totals
    total_row = len(summary_data) + 5
    ws_sum.cell(row=total_row, column=1, value="GESAMT").font = bold_font
    ws_sum.cell(row=total_row, column=2, value=format_duration(dept_soll)).font = bold_font
    ws_sum.cell(row=total_row, column=3, value=format_duration(dept_ist)).font = bold_font
    delta_cell = ws_sum.cell(row=total_row, column=4, value=format_duration(dept_ist - dept_soll))
    delta_cell.font = bold_font

    ws_sum.column_dimensions["A"].width = 25
    for c in ["B", "C", "D", "E"]:
        ws_sum.column_dimensions[c].width = 15

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
